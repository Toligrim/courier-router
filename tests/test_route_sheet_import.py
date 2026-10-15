"""Printed route-sheet layout, with synthetic contact and address data."""
import csv
from io import BytesIO

import pytest
from openpyxl import Workbook

from courier_router import pdf_converter, web
from courier_router.parsing import read_table


HEADER = ["#", "Тип", "Заказ", "Адрес", "Район", "Телефон", "Сумма", "Оплата", "Время", "Комментарий", "Отметка"]
ROWS = [
    [1, "Забор", "#00503", "Тестовая ул., д 17 к 2, кв. 237", "Невский", "+7 (900) 000-00-01", None, "Бесплатно / гарантия", "17:00-20:30", None, None],
    [2, "Забор", "#00490", "Тестовая ул., д 115 к 1, кв. 66", "Невский", "+7 (900) 000-00-02", None, "—", "17:00-20:30", "в любой день, кроме понедельника", None],
    [3, "Забор", "#00479", "Тестовый бульвар, д 3 кв 161", "Пушкинский", "+7 (900) 000-00-03", None, "—", "17:00-20:30", None, None],
    [4, "Забор", "#00458", "Тестовое шоссе, д 37 к 1, кв. 339", "Красногвардейский", "+7 (900) 000-00-04", None, "—", "17:00-20:30", "9 этаж 7 парадная", None],
    [5, "Отвоз", "#00339", "Тестовая ул., д 42", "Невский", "+7 (900) 000-00-05", 2900, "Наличные", "17:00-20:30", None, None],
]


def workbook_bytes(rows=None, header=None, declared=5):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Маршрутный лист — Тестовый курьер", None, None, None, None, "Название партнёра"])
    sheet.merge_cells("A1:E1")
    sheet.merge_cells("F1:K1")
    sheet.append([f"четверг, 15 октября 2026 г. · всего {declared} выездов"])
    sheet.merge_cells("A2:K2")
    sheet.append([])
    sheet.append(HEADER if header is None else header)
    for row in ROWS if rows is None else rows:
        sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def test_partner_workbook_preserves_all_route_fields(tmp_path):
    path = tmp_path / "route.xlsx"
    path.write_bytes(workbook_bytes())
    stops = read_table(path, default_service_min=12)

    assert [stop.source_row for stop in stops] == [5, 6, 7, 8, 9]
    assert [stop.order_no for stop in stops] == [503, 490, 479, 458, 339]
    assert [stop.operation.value for stop in stops] == ["pickup"] * 4 + ["delivery"]
    for stop, row in zip(stops, ROWS):
        assert stop.address_raw == row[3]
        assert stop.district == row[4]
        assert stop.phone.endswith(row[5][-12:])
        assert stop.access == ""
        assert stop.window.raw == "17:00-20:30"
        assert (stop.window.start_min, stop.window.end_min) == (1020, 1230)
        assert stop.service_min == 12
    assert stops[0].payment.amount_rub == 0
    assert stops[0].payment.method == "free"
    assert stops[0].payment.raw == "Бесплатно / гарантия"
    assert stops[1].payment.amount_rub is None
    assert stops[1].payment.raw == ""
    assert stops[1].comment == "в любой день, кроме понедельника"
    assert stops[3].comment == "9 этаж 7 парадная"
    assert stops[4].payment.amount_rub == 2900
    assert stops[4].payment.method == "cash"


def test_partner_csv_uses_headers_instead_of_column_positions(tmp_path):
    path = tmp_path / "route.csv"
    # Same headings in reverse order, no title rows; also keep the optional mark.
    rows = [row.copy() for row in ROWS]
    rows[0][-1] = "Позвонить заранее"
    rows[1][8] = "—"
    rows[2][6:8] = [0, "Перевод"]
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        csv.writer(stream, delimiter=";").writerows([list(reversed(row)) for row in [HEADER, *rows]])

    stops = read_table(path)
    assert stops[0].order_no == 503
    assert stops[0].comment == "Отметка: Позвонить заранее"
    assert stops[1].window is None
    assert stops[2].payment.amount_rub == 0
    assert stops[2].payment.method == "transfer"
    assert stops[4].payment.amount_rub == 2900


@pytest.mark.parametrize("column,value,message", [
    (1, "Неизвестно", "Строка 6, колонка B"),
    (8, "25:00-26:00", "Строка 6, колонка I"),
    (2, "#bad", "Строка 6: некорректный номер заказа"),
    (0, 3, "Строка 6: нарушена последовательность"),
    (3, "", "Строка 6: пустой адрес"),
])
def test_partner_errors_report_original_locations(tmp_path, column, value, message):
    rows = [row.copy() for row in ROWS]
    rows[1][column] = value
    path = tmp_path / "bad.xlsx"
    path.write_bytes(workbook_bytes(rows))
    with pytest.raises(ValueError, match=message):
        read_table(path)


def test_partner_count_catches_missing_final_order(tmp_path):
    path = tmp_path / "truncated.xlsx"
    path.write_bytes(workbook_bytes(ROWS[:-1]))
    with pytest.raises(ValueError, match="Количество заказов не совпало"):
        read_table(path)


@pytest.mark.parametrize("header,message", [
    ([value if value != "Время" else "???" for value in HEADER], "отсутствуют колонки: время"),
    (HEADER + ["Тип"], "повторяющиеся заголовки"),
    (["Непонятная таблица"], "не найдена строка заголовков"),
])
def test_partner_rejects_invalid_headers(tmp_path, header, message):
    path = tmp_path / "bad.xlsx"
    path.write_bytes(workbook_bytes(header=header))
    with pytest.raises(ValueError, match=message):
        read_table(path)


def test_pdf_is_supported_by_shared_parser(monkeypatch, tmp_path):
    source = tmp_path / "route.PDF"
    source.write_bytes(b"%PDF-test")
    monkeypatch.setattr(pdf_converter, "_ocr_pages", lambda *args: (
        "Маршрутный лист — всего 1 выезд\n"
        "# Тип Заказ Адрес Район Телефон Сумма Оплата Время Комментарий Отметка\n"
        "1 Отвоз #00339 Тестовая ул., д 42 Невский +7 (900) 000-00-05 2900Р Наличные 17:00-20:30\n"
    ))
    stops = read_table(source, default_service_min=15)
    assert len(stops) == 1
    assert stops[0].order_no == 339
    assert stops[0].service_min == 15
    assert stops[0].payment.amount_rub == 2900
    assert stops[0].payment.method == "cash"
    assert stops[0].window.raw == "17:00-20:30"
    assert not source.with_suffix(".xlsx").exists()


def test_web_partner_xlsx_reaches_planner(monkeypatch, tmp_path):
    from test_web import _client_logged_in

    client = _client_logged_in(monkeypatch, tmp_path)
    parsed = []

    def plan(args):
        parsed.extend(read_table(args.xlsx))

    monkeypatch.setattr(web, "cmd_plan", plan)
    response = client.post(
        "/routes", data={"day": "2026-10-15", "depart": "17:00", "end": "open"},
        files={"table": ("Маршрутный лист.xlsx", workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        follow_redirects=False,
    )
    assert response.status_code == 303
    assert [stop.order_no for stop in parsed] == [503, 490, 479, 458, 339]
    assert parsed[-1].payment.amount_rub == 2900
