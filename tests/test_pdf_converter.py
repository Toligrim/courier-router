import pytest

from openpyxl import load_workbook

from courier_router.parsing import read_table
from courier_router.pdf_converter import (
    PdfConversionError,
    _OCR_SLOT,
    convert_pdf_to_xlsx,
    parse_ocr_text,
    write_xlsx,
)


OCR_22 = """Маршрутный лист — Носов Анатолий Коврокот — всего 6 выездов
# Тип Заказ Адрес Район Телефон Сумма Оплата Время Комментарий Отметка
1 Забор #00442 г Санкт-Петербург, ул Брянцева, д 14, кв. 63 Калининский +7 (911) 815-60-15 — 17:00-20:30
2 Забор #00436 г Санкт-Петербург, поселок Парголово, ул Выборгский +7 (921) 750-10-23 — 17:00-20:30
Фёдора Абрамова, д 15, кв. 217
3 Забор #00435 г Санкт-Петербург, Богатырский пр-кт, д Зк1, кв. Приморский +7 (911) 813-09-46 — 17:00-20:30
13
4 Забор #00428 г Санкт-Петербург, ул Парашютная, д 22 к1, кв. Приморский +7 (961) 801-90-07 — 17:00-20:30
271
5 Забор #00424 г Санкт-Петербург, пр-кт Энгельса, д 93, кв. 226 Выборгский +7 (900) 623-70-20 — 17:00-20:30
6 Отвоз #00332 г Санкт-Петербург, пр-кт Науки, д 17к2 Калининский +7 (931) 234-14-35 2300Р Наличные 17:00-20:30
"""

OCR_23 = """Маршрутный лист — Носов Анатолий Коврокот — всего 4 выездов
# Тип Заказ Адрес Район Телефон Сумма Оплата Время Комментарий Отметка
1 Забор #00446 г Санкт-Петербург, Театральная пл, д 16, кв. 16 Адмиралтейский +7 (911) 936-97-47 — — 19:30-20:30
2 Отвоз #00399 г Санкт-Петербург, Театральная пл, д 16, кв. 16 Адмиралтейский +7 (911) 936-97-47 3300Р — 19:30 -20:30
3 Отвоз #00430 г Санкт-Петербург, Крестовский пр-кт, д 10 Петроградский +7 (921) 750-69-53 5200Р Наличные 17:00-20:30
4 Отвоз #00415 г Санкт-Петербург, ул Можайская, д 3-5, кв. 19 Адмиралтейский +7 (981) 185-59-20 2200Р — 17:00-20:30
"""


def test_parse_partner_route_sheets():
    rows22 = parse_ocr_text(OCR_22)
    rows23 = parse_ocr_text(OCR_23)

    assert len(rows22) == 6
    assert rows22[1].address.endswith("ул Фёдора Абрамова, д 15, кв. 217")
    assert rows22[2].address.endswith("д 3 к 1, кв. 13")
    assert rows22[5].payment == "2300 / Наличные"
    assert len(rows23) == 4
    assert rows23[1].window == "19:30-20:30"
    assert rows23[2].payment == "5200 / Наличные"


def test_xlsx_output_is_accepted_by_route_parser(tmp_path):
    output = write_xlsx(parse_ocr_text(OCR_23), tmp_path / "route.xlsx")
    assert output.exists()
    assert load_workbook(output, read_only=True).active.max_column == 9

    stops = read_table(output)
    assert len(stops) == 4
    assert stops[0].order_no == 446
    assert stops[1].payment.amount_rub == 3300
    assert stops[2].payment.method == "cash"


def test_parser_rejects_missing_order_row():
    broken = OCR_23.replace(
        "2 Отвоз #00399", "7 Отвоз #00399"
    )
    try:
        parse_ocr_text(broken)
    except ValueError as exc:
        assert "последовательность" in str(exc)
    else:
        raise AssertionError("missing order sequence must be rejected")


def test_parser_rejects_declared_count_mismatch():
    broken = OCR_23.replace("всего 4 выездов", "всего 5 выездов")
    try:
        parse_ocr_text(broken)
    except ValueError as exc:
        assert "Количество заказов" in str(exc)
    else:
        raise AssertionError("declared count mismatch must be rejected")


def test_parser_rejects_unknown_continuation_instead_of_changing_address():
    broken = OCR_23.replace(
        "2 Отвоз #00399", "OCR_GARBAGE\n2 Отвоз #00399"
    )
    try:
        parse_ocr_text(broken)
    except ValueError as exc:
        assert "неизвестное продолжение" in str(exc)
    else:
        raise AssertionError("unknown OCR continuation must be rejected")

    broken_short = OCR_23.replace(
        "2 Отвоз #00399", "к двери\n2 Отвоз #00399"
    )
    try:
        parse_ocr_text(broken_short)
    except ValueError as exc:
        assert "неизвестное продолжение" in str(exc)
    else:
        raise AssertionError("non-address continuation must be rejected")


def test_parser_rejects_hash_continuation_instead_of_skipping_it():
    broken = OCR_23.replace(
        "2 Отвоз #00399", "# OCR_GARBAGE\n2 Отвоз #00399"
    )
    try:
        parse_ocr_text(broken)
    except ValueError as exc:
        assert "неизвестное продолжение" in str(exc)
    else:
        raise AssertionError("hash continuation must be rejected")


def test_parser_rejects_unknown_text_before_first_order():
    broken = OCR_23.replace(
        "1 Забор #00446", "OCR_GARBAGE\n1 Забор #00446"
    )
    try:
        parse_ocr_text(broken)
    except PdfConversionError as exc:
        assert "перед первой строкой" in str(exc)
    else:
        raise AssertionError("unknown preamble must be rejected")


def test_parser_rejects_duplicate_order_numbers():
    broken = OCR_23.replace("#00430", "#00446")
    try:
        parse_ocr_text(broken)
    except PdfConversionError as exc:
        assert "повторяющиеся" in str(exc)
    else:
        raise AssertionError("duplicate order numbers must be rejected")


def test_parser_rejects_unknown_payment_value():
    broken = OCR_23.replace("3300Р —", "3300Р Сомнительно")
    try:
        parse_ocr_text(broken)
    except PdfConversionError as exc:
        assert "значение оплаты" in str(exc)
    else:
        raise AssertionError("unknown payment value must be rejected")


def test_converter_rejects_parallel_ocr_job(tmp_path):
    assert _OCR_SLOT.acquire(blocking=False)
    try:
        try:
            convert_pdf_to_xlsx(tmp_path / "missing.pdf", tmp_path / "route.xlsx")
        except PdfConversionError as exc:
            assert "другая PDF-конвертация" in str(exc)
        else:
            raise AssertionError("parallel conversion must be rejected")
    finally:
        _OCR_SLOT.release()


@pytest.mark.parametrize("angle", [0, 90, 180, 270])
def test_orientation_is_detected_on_preview_before_full_resolution_ocr(monkeypatch, tmp_path, angle):
    from types import SimpleNamespace
    from pathlib import Path
    from PIL import Image
    from courier_router import pdf_converter as converter

    source = tmp_path / "scan.pdf"
    source.write_bytes(b"%PDF-test")
    monkeypatch.setattr(converter.shutil, "which", lambda name: name)

    def run(command, **kwargs):
        if command[0] == "pdfinfo":
            return SimpleNamespace(stdout="Pages: 1\nPage size: 595 x 842 pts\n")
        assert command[0] == "pdftoppm"
        Image.new("RGB", (2400, 3400), "white").save(command[-1] + "-1.png")
        return SimpleNamespace(stdout="")

    monkeypatch.setattr(converter.subprocess, "run", run)
    calls = []
    preview_count = angle // 90 + 1

    def ocr(path, command, language, timeout):
        with Image.open(path) as image:
            calls.append(image.size)
        assert 0 < timeout <= converter._OCR_PROCESS_TIMEOUT
        if len(calls) < preview_count:
            return "unreadable sideways text"
        if len(calls) == preview_count:
            return "Маршрутный лист Тип Заказ Адрес Район Телефон Время"
        assert Path(path).name.endswith("-oriented.png")
        return OCR_23

    monkeypatch.setattr(converter, "_run_ocr", ocr)
    text = converter._ocr_pages(source, "rus+eng", "tesseract")
    assert len(converter.parse_ocr_text(text)) == 4
    assert len(calls) == preview_count + 1
    assert all(max(size) <= 2000 for size in calls[:-1])
    assert calls[-1] == ((3400, 2400) if angle in (90, 270) else (2400, 3400))


def test_ocr_limits_thread_count(monkeypatch, tmp_path):
    from types import SimpleNamespace
    from courier_router import pdf_converter as converter

    calls = []

    def run(command, **kwargs):
        calls.append((command, kwargs))
        return SimpleNamespace(stdout="recognized")

    monkeypatch.setattr(converter.subprocess, "run", run)
    assert converter._run_ocr(tmp_path / "page.png", "tesseract", "rus+eng", 5) == "recognized"
    assert calls[0][1]["env"]["OMP_THREAD_LIMIT"] == "1"


def test_render_timeout_is_reported_as_conversion_error(monkeypatch, tmp_path):
    from types import SimpleNamespace
    from courier_router import pdf_converter as converter

    source = tmp_path / "scan.pdf"
    source.write_bytes(b"%PDF-test")
    monkeypatch.setattr(converter.shutil, "which", lambda name: name)

    def run(command, **kwargs):
        if command[0] == "pdfinfo":
            return SimpleNamespace(stdout="Pages: 1\nPage size: 595 x 842 pts\n")
        raise converter.subprocess.TimeoutExpired(command, 30)

    monkeypatch.setattr(converter.subprocess, "run", run)
    with pytest.raises(PdfConversionError, match="Рендеринг PDF превысил лимит времени"):
        converter._ocr_pages(source, "rus+eng", "tesseract")


def _two_page_ocr():
    lines = OCR_23.splitlines(keepends=True)
    boundary = next(i for i, line in enumerate(lines) if line.startswith("3 Отвоз"))
    return "".join(lines[:boundary]), "".join(lines[:2] + lines[boundary:])


def test_multiple_pages_keep_repeated_headings_out_of_addresses():
    first, second = _two_page_ocr()
    rows = parse_ocr_text(first + "\f" + second)
    assert rows == parse_ocr_text(OCR_23)


def test_multiple_pages_still_validate_sequence_and_total():
    first, second = _two_page_ocr()
    with pytest.raises(PdfConversionError, match="последовательность"):
        parse_ocr_text(first + "\f" + second.replace("3 Отвоз", "7 Отвоз"))
    with pytest.raises(PdfConversionError, match="Количество заказов"):
        parse_ocr_text(first + "\f" + second.replace("всего 4 выездов", "всего 5 выездов"))


def test_multiple_pages_require_header_on_each_page():
    first, second = _two_page_ocr()
    second = "\n".join(second.splitlines()[2:])
    with pytest.raises(PdfConversionError, match="не найдена строка заголовков"):
        parse_ocr_text(first + "\f" + second)


def test_ocr_accepts_numbered_pdfinfo_page_sizes(monkeypatch, tmp_path):
    from types import SimpleNamespace
    from pathlib import Path
    from PIL import Image
    from courier_router import pdf_converter as converter

    source = tmp_path / "two-pages.pdf"
    source.write_bytes(b"%PDF-test")
    monkeypatch.setattr(converter.shutil, "which", lambda name: name)
    checked_pages = []

    def run(command, **kwargs):
        if command[0] == "pdfinfo":
            if "-f" in command:
                checked_pages.append(int(command[command.index("-f") + 1]))
                return SimpleNamespace(stdout="Pages: 2\nPage    2 size: 595 x 842 pts (A4)\n")
            return SimpleNamespace(stdout="Pages: 2\nPage size: 595 x 842 pts (A4)\n")
        assert command[0] == "pdftoppm"
        for page in (1, 2):
            Image.new("RGB", (300, 400), "white").save(command[-1] + f"-{page}.png")
        return SimpleNamespace(stdout="")

    monkeypatch.setattr(converter.subprocess, "run", run)
    first, second = _two_page_ocr()

    def ocr(path, *args):
        if Path(path).name.endswith("-preview.png"):
            return "Маршрутный лист Тип Заказ Адрес Район Телефон Время"
        return first if Path(path).name.startswith("page-1-") else second

    monkeypatch.setattr(converter, "_run_ocr", ocr)
    text = converter._ocr_pages(source, "rus+eng", "tesseract")
    assert checked_pages == [2]
    assert "\f" in text
    assert parse_ocr_text(text) == parse_ocr_text(OCR_23)
