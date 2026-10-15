from __future__ import annotations
import csv
import re
import unicodedata
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from .domain import Operation, Payment, Stop, TimeWindow

_CLOCK_TOKEN = r"\d{1,2}\s*[:.\-–—−]\s*\d{1,2}(?:\s*:\s*\d{1,2})?"
WINDOW_RE = re.compile(
    rf"(?<!\d)(?P<start>{_CLOCK_TOKEN})"
    rf"(?P<separator>\s*(?:[-–—−]|\bдо\b)\s*|\s+)"
    rf"(?P<end>{_CLOCK_TOKEN})(?!\d)",
    re.IGNORECASE,
)
HOUR_WINDOW_RE = re.compile(
    r"(?<![\d:.\-–—−])(?P<h1>\d{1,2})\s*(?:[-–—−]|\bдо\b)\s*"
    r"(?P<h2>\d{1,2})(?!\s*[\d:.\-–—−])",
    re.IGNORECASE,
)
COMPACT_WINDOW_RE = re.compile(
    r"(?P<h1>\d{1,2})\s*[-–—−]\s*(?P<m1>\d{2})\s*[-–—−]\s*"
    r"(?P<h2>\d{1,2})\s*[-–—−]\s*(?P<m2>\d{2})"
)
CLOCK_RE = re.compile(
    r"(?<!\d)(?P<h>\d{1,2})\s*[:.\-–—−]\s*(?P<m>\d{1,2})"
    r"(?:\s*:\s*(?P<s>\d{1,2}))?(?!\d)"
)
EXTRA_NUMBER_RE = re.compile(r"\d+")
MONEY_RE = re.compile(r"(?P<amount>\d[\d\s]*)")
OPERATION_ALIASES = {
    "забор": Operation.PICKUP,
    "забрать": Operation.PICKUP,
    "заб": Operation.PICKUP,
    "отвоз": Operation.DELIVERY,
    "отвезти": Operation.DELIVERY,
    "отв": Operation.DELIVERY,
    "доставка": Operation.DELIVERY,
    "доставить": Operation.DELIVERY,
    "дост": Operation.DELIVERY,
}


def _normalize_operation(value: object) -> tuple[Operation, str]:
    """Recognize an operation and retain an optional clarification for the courier."""
    text = unicodedata.normalize("NFKC", str(value or ""))
    text = " ".join(text.replace("\ufeff", "").replace("\u200b", "").split())
    words = re.findall(r"\w+", text.casefold())
    operations = {OPERATION_ALIASES[word] for word in words if word in OPERATION_ALIASES}
    if len(operations) > 1:
        raise ValueError(
            f"Неоднозначный тип операции: {value!r}. "
            "Укажите одну операцию: Забор, Отвоз или Доставка"
        )
    if not operations or any(word not in OPERATION_ALIASES and word != "уточнить" for word in words):
        raise ValueError(
            f"Неизвестный тип операции: {value!r}. "
            "Ожидается Забор, Отвоз или Доставка; допустима пометка «Уточнить»"
        )
    note = f"Тип операции в таблице: {text}" if "уточнить" in words else ""
    return operations.pop(), note


def parse_operation(value: object) -> Operation:
    return _normalize_operation(value)[0]


def _can_be_hour_window(value: str) -> bool:
    match = HOUR_WINDOW_RE.fullmatch(value.strip())
    if not match:
        return False
    start, end = int(match["h1"]), int(match["h2"])
    return 0 <= start <= 23 and start <= end <= 24


def parse_window(value: object) -> TimeWindow | None:
    """Parse a window and store it as canonical ``HH:MM-HH:MM`` text."""
    s = unicodedata.normalize("NFKC", str(value or "")).replace("\xa0", " ").strip()
    if not s:
        return None

    # Some spreadsheets use dashes both inside each clock and between clocks,
    # e.g. ``20-00-20-30``. Normalize this compact form only when it cannot
    # also mean two separate whole-hour windows such as ``17-18-19-20``.
    compact = COMPACT_WINDOW_RE.fullmatch(s)
    if compact:
        h1, m1, h2, m2 = (int(compact[name]) for name in ("h1", "m1", "h2", "m2"))
        two_hour_windows = (
            0 <= h1 <= m1 <= 24
            and 0 <= h2 <= m2 <= 24
        )
        if not two_hour_windows:
            s = f"{h1}:{m1:02d}-{h2}:{m2:02d}"

    m = WINDOW_RE.search(s)
    if m:
        start_token, end_token = m["start"], m["end"]
        separator = m["separator"]
    else:
        # Частый ручной формат: «17-18» означает «17:00-18:00».
        m = HOUR_WINDOW_RE.search(s)
        if not m:
            raise ValueError(f"Не удалось разобрать временное окно: {s!r}")
        start_token = f"{m['h1']}:00"
        end_token = f"{m['h2']}:00"
        separator = ""

    match_start, match_end = m.span()
    dash_chars = "-–—−"
    start_has_dash = any(char in dash_chars for char in start_token)
    end_has_dash = any(char in dash_chars for char in end_token)
    separator_is_dash = bool(separator.strip()) and separator.strip() in dash_chars
    ambiguous_dash_clock = (
        separator_is_dash
        and (
            (start_has_dash and end_has_dash)
            or ((start_has_dash or end_has_dash) and not any(char.isspace() for char in separator))
        )
    )
    ambiguous_hour_windows = (
        not separator.strip()
        and (
            _can_be_hour_window(start_token)
            or _can_be_hour_window(end_token)
        )
    )
    if (
        (match_start and s[match_start - 1].isalnum())
        or (match_end < len(s) and s[match_end].isalnum())
        or ambiguous_dash_clock
        or ambiguous_hour_windows
        or any(
            clock.start() < match_start or clock.end() > match_end
            for clock in CLOCK_RE.finditer(s)
        )
        or any(
            number.start() < match_start or number.end() > match_end
            for number in EXTRA_NUMBER_RE.finditer(s)
        )
    ):
        raise ValueError(f"Временное окно содержит лишние значения времени: {s!r}")

    start_match = CLOCK_RE.fullmatch(start_token.strip())
    end_match = CLOCK_RE.fullmatch(end_token.strip())
    if not start_match or not end_match:
        raise ValueError(f"Не удалось разобрать временное окно: {s!r}")

    h1, m1 = int(start_match["h"]), int(start_match["m"])
    h2, m2 = int(end_match["h"]), int(end_match["m"])
    seconds = [start_match["s"], end_match["s"]]
    if any(item is not None and int(item) != 0 for item in seconds):
        raise ValueError(f"Временное окно должно быть указано с точностью до минуты: {s!r}")
    if not (0 <= h1 <= 23 and 0 <= m1 <= 59):
        raise ValueError(f"Некорректное начало временного окна: {s!r}")
    if not (0 <= h2 <= 24 and 0 <= m2 <= 59) or (h2 == 24 and m2 != 0):
        raise ValueError(f"Некорректный конец временного окна: {s!r}")
    a = h1 * 60 + m1
    b = h2 * 60 + m2
    if b < a:
        raise ValueError(f"Окно заканчивается раньше начала: {s!r}")
    return TimeWindow(a, b, f"{h1:02d}:{m1:02d}-{h2:02d}:{m2:02d}")


def parse_payment(value: object) -> Payment:
    s = str(value or "").strip()
    if not s:
        return Payment(raw="")
    low = s.lower()
    if "бесплат" in low:
        return Payment(amount_rub=0, method="free", raw=s)
    m = MONEY_RE.search(low.replace("\xa0", " "))
    amount = int(re.sub(r"\s+", "", m["amount"])) if m else None
    if "безнал" in low or "перев" in low or "кар" in low:
        method = "transfer"
    elif "нал" in low:
        method = "cash"
    else:
        method = "unknown"
    return Payment(amount_rub=amount, method=method, raw=s)


_PHONE_SPLIT = re.compile(r"\s*(?:/|,|;|\bили\b)\s*", re.I)


def _format_ru_phone(digits: str) -> str | None:
    """Цифры российского номера → «8 (XXX) XXX-XX-XX» (+ « доб. N», если есть). Иначе None."""
    ext = ""
    if len(digits) > 11 and digits[0] in "78":
        digits, ext = digits[:11], digits[11:]
    if len(digits) == 11 and digits[0] in "78":
        digits = "8" + digits[1:]
    elif len(digits) == 10 and digits[0] == "9":
        digits = "8" + digits            # номер без кода страны
    else:
        return None
    core = f"{digits[0]} ({digits[1:4]}) {digits[4:7]}-{digits[7:9]}-{digits[9:11]}"
    return f"{core} доб. {ext}" if ext else core


def normalize_phone(value: object) -> str:
    """Единый вид номеров, чтобы по ним нормально звонилось: ведущие +7/7/8 → 8,
    формат «8 (921) 123-45-67». Несколько номеров в ячейке разделяются « / »."""
    s = str(value or "").strip()
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    if not s:
        return ""
    out: list[str] = []
    for token in _PHONE_SPLIT.split(s):
        token = token.strip()
        if not token:
            continue
        out.append(_format_ru_phone(re.sub(r"\D", "", token)) or token)
    return " / ".join(dict.fromkeys(out))


def phone_dial_digits(display: str) -> str:
    """Цифры первого номера для ссылки tel: (например «89211234567»)."""
    first = (display or "").split("/")[0].split("доб")[0]
    digits = re.sub(r"\D", "", first)
    if len(digits) >= 11 and digits[0] in "78":
        digits = "8" + digits[1:11]
    elif len(digits) == 10 and digits[0] == "9":
        digits = "8" + digits
    return digits


_ROUTE_COLUMNS = ("тип", "заказ", "телефон", "район", "адрес", None, "время", "оплата", "комментарий")
_REQUIRED_ROUTE_COLUMNS = {"тип", "заказ", "адрес", "район", "телефон", "сумма", "оплата", "время"}
_DECLARED_COUNT_RE = re.compile(r"\bвсего\s+(\d+)\s+выезд", re.IGNORECASE)


def _cell_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _route_cell(value: object) -> str:
    text = _cell_text(value)
    return "" if text in {"—", "–", "-"} else text


def _table_rows(rows):
    """Map the printed partner layout to the original A-I layout, keeping locations."""
    columns = None
    preamble = False
    data_started = False
    declared_counts = []
    count = 0
    for row_no, row in enumerate(rows, 1):
        if not any(_cell_text(value) for value in row):
            continue
        if not data_started:
            labels = [unicodedata.normalize("NFKC", _cell_text(v)).casefold() for v in row]
            if {"тип", "заказ", "адрес"}.issubset(labels):
                missing = _REQUIRED_ROUTE_COLUMNS.difference(labels)
                if missing:
                    raise ValueError(f"Строка {row_no}: отсутствуют колонки: {', '.join(sorted(missing))}")
                if len(set(labels) - {""}) != len([label for label in labels if label]):
                    raise ValueError(f"Строка {row_no}: повторяющиеся заголовки колонок")
                columns = {label: i for i, label in enumerate(labels) if label}
                data_started = True
                continue
            if preamble or any(label.startswith("маршрутный лист") for label in labels):
                preamble = True
                declared_counts.extend(int(m[1]) for value in row for m in _DECLARED_COUNT_RE.finditer(_cell_text(value)))
                continue
            data_started = True

        if columns is None:
            yield row_no, list(row[:9]) + [None] * max(0, 9 - len(row)), "A", "G"
            continue

        def cell(label):
            index = columns.get(label)
            return row[index] if index is not None and index < len(row) else None

        count += 1
        if "#" in columns and _cell_text(cell("#")) != str(count):
            raise ValueError(f"Строка {row_no}: нарушена последовательность, ожидался номер {count}")
        vals = [cell(label) if label else "" for label in _ROUTE_COLUMNS]
        order = _cell_text(vals[1]).removeprefix("#").strip()
        if not re.fullmatch(r"\d+", order):
            raise ValueError(f"Строка {row_no}: некорректный номер заказа: {vals[1]!r}")
        vals[1] = int(order)
        vals[6] = _route_cell(vals[6])
        amount = _route_cell(cell("сумма"))
        vals[7] = " / ".join(filter(None, (amount, _route_cell(vals[7]))))
        mark = _route_cell(cell("отметка"))
        vals[8] = "; ".join(filter(None, (_route_cell(vals[8]), f"Отметка: {mark}" if mark else "")))
        yield row_no, vals, get_column_letter(columns["тип"] + 1), get_column_letter(columns["время"] + 1)

    if preamble and columns is None:
        raise ValueError("В маршрутном листе не найдена строка заголовков таблицы")
    if columns is not None and any(declared != count for declared in declared_counts):
        raise ValueError(f"Количество заказов не совпало: в заголовке {declared_counts[0]}, прочитано {count}")


def _rows_to_stops(rows, default_service_min: int = 10) -> list[Stop]:
    stops: list[Stop] = []
    for row_no, vals, operation_column, window_column in _table_rows(rows):
        if not any(v not in (None, "") for v in vals[:8]):
            continue
        try:
            op, operation_note = _normalize_operation(vals[0])
        except ValueError as e:
            raise ValueError(f"Строка {row_no}, колонка {operation_column}: {e}") from e
        try:
            order_no = int(vals[1]) if isinstance(vals[1], int) else int(float(vals[1]))
        except Exception as e:
            raise ValueError(f"Строка {row_no}: {e}") from e
        address = str(vals[4] or "").strip()
        if not address:
            raise ValueError(f"Строка {row_no}: пустой адрес")
        try:
            window = parse_window(vals[6])
        except ValueError as e:
            raise ValueError(f"Строка {row_no}, колонка {window_column}: {e}") from e
        stops.append(Stop(
            source_row=row_no,
            operation=op,
            order_no=order_no,
            phone=normalize_phone(vals[2]),
            district=str(vals[3] or "").strip(),
            address_raw=address,
            access=str(vals[5] or "").strip(),
            window=window,
            payment=parse_payment(vals[7]),
            comment="; ".join(filter(None, (operation_note, str(vals[8] or "").strip()))),
            service_min=default_service_min,
        ))
    if not stops:
        raise ValueError("В таблице не найдено ни одной точки")
    return stops


def read_excel(path: str | Path, default_service_min: int = 10) -> list[Stop]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return _rows_to_stops(wb.active.iter_rows(values_only=True), default_service_min)
    finally:
        wb.close()


def read_csv(path: str | Path, default_service_min: int = 10) -> list[Stop]:
    path = Path(path)
    raw = path.read_bytes()
    text = None
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("CSV должен быть в UTF-8 или Windows-1251")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
        reader = csv.reader(text.splitlines(), dialect)
    except csv.Error:
        reader = csv.reader(text.splitlines(), delimiter=";")
    return _rows_to_stops(reader, default_service_min)


def read_table(path: str | Path, default_service_min: int = 10) -> list[Stop]:
    suffix = Path(path).suffix.lower()
    if suffix == ".xlsx":
        return read_excel(path, default_service_min)
    if suffix == ".csv":
        return read_csv(path, default_service_min)
    if suffix == ".pdf":
        from .pdf_converter import read_pdf_rows

        return _rows_to_stops((row.as_xlsx_row() for row in read_pdf_rows(path)), default_service_min)
    raise ValueError("Поддерживаются файлы .xlsx, .csv и .pdf")
